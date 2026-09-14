from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT='/Users/mmiao/Desktop/液体智能网页/液态智能网页功能说明.docx'
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
styles=doc.styles; styles['Normal'].font.name='PingFang SC'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'PingFang SC'); styles['Normal'].font.size=Pt(10.5); styles['Normal'].font.color.rgb=RGBColor(45,45,55)
for name,size,color in [('Title',27,'11111E'),('Heading 1',17,'11111E'),('Heading 2',12,'E35D45')]:
 s=styles[name]; s.font.name='PingFang SC'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'PingFang SC'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
def shade(cell,color):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); tcPr.append(shd)
def table(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'
 for c,x in zip(t.rows[0].cells,headers): c.text=x; shade(c,'20203A')
 for r in rows:
  cells=t.add_row().cells
  for c,x in zip(cells,r): c.text=x
 return t
doc.add_paragraph('液态智能网页功能说明',style='Title'); p=doc.add_paragraph('产品功能、交互流程、部署状态与后续边界'); p.runs[0].font.color.rgb=RGBColor(227,93,69); doc.add_paragraph('文档版本 1.0  ·  2026年9月',style='Subtitle')
doc.add_heading('文档概述',1); doc.add_paragraph('液态智能网页是一套“从自然语言需求到可操作网页”的生成系统。用户只需输入一次完整 Query，系统会先理解目标、约束和事实，再生成一份统一的任务答案，并以七种不同的网页形态呈现。它适用于商品推荐、方案比较、学习规划、旅行安排、项目管理、预算计算和数据分析等任务。')
doc.add_paragraph('当前版本已经完成可公开演示的核心闭环，并部署在 Render。它更适合作为产品原型和邀请制试用版本；涉及价格、库存或专业结论的内容，仍应在行动前再次核实。')
doc.add_heading('核心功能',1)
for h,txt in [('统一 Query 输入','支持输入目标、预算、偏好、已有资料和限制条件，最大长度为 6000 字符。'),('需求与约束理解','自动提取用户目标、硬性条件、用户提供信息、待核实事实、假设和预期交付结果。'),('统一答案与七种形态','同一份事实、比较结果和最终结论会被复用到手账决策本、双区分析台、专题编辑部、ICE 速览报、留白阅读页、任务工作台和数据驾驶舱。'),('任务型内容生成','根据任务选择思考路径、步骤、卡片、表格、时间线、对比、计算器、图表和清单等组件。'),('品牌与选购入口','购买类需求至少列出三个候选项，并提供京东、淘宝或官网搜索入口。链接由系统安全生成，不使用模型随意编造的商品详情页地址。'),('事实状态展示','内容可以标记为用户提供、已验证或待核实，减少不同页面之间的事实冲突。'),('可操作交互','支持点击比较表行、调节计算器滑块、勾选清单、使用模块导航和切换七种形态。'),('状态恢复','浏览器会保存 Query、已生成页面、计算器数值和清单勾选状态；服务端会话有效期为两小时。'),('错误与服务保护','支持模型超时提示、临时拥堵自动重试、API Key 检查、健康检查、单 IP 请求限制和基础安全响应头。')]:
 doc.add_heading(h,2); doc.add_paragraph(txt)
doc.add_heading('用户使用流程',1); table(['阶段','用户操作','系统结果'],[('1 输入','填写完整 Query 并提交','理解目标、提取约束和整理共享事实'),('2 选择','浏览七张网页形态卡片','七种形态共享同一份答案，不再各自产生互相矛盾的结论'),('3 展开','点击任意卡片','显示生成进度，生成完成后进入完整任务页面'),('4 操作','查看证据、比较候选、调节参数或勾选清单','页面状态即时更新，并保存到当前浏览器'),('5 跳转','点击品牌或平台入口','在新标签页打开对应的站内搜索页面')])
doc.add_heading('页面形态',1); table(['形态','主要特点','适合场景'],[('手账决策本','纸张纹理、标签、推理路径、彩色结论卡','商品推荐、生活决策、个人规划'),('双区分析台','终端色彩、指标和结构化数据','技术分析、参数比较、预算计算'),('专题编辑部','杂志式排版和分栏内容','专题介绍、方案评审、内容阅读'),('ICE 速览报','高对比标题和快速扫描结构','快速了解结论和重点'),('留白阅读页','低干扰、长阅读、重点突出','解释、学习和报告阅读'),('任务工作台','任务卡、状态和可勾选事项','项目计划、执行清单、日程安排'),('数据驾驶舱','指标卡、图表和高密度信息','数据分析和监控类任务')])
doc.add_heading('部署与访问',1); doc.add_paragraph('当前网站已经部署到 Render，可通过以下地址访问：'); p=doc.add_paragraph('https://liquid-intelligent-web.onrender.com'); p.runs[0].font.color.rgb=RGBColor(35,95,180); p.runs[0].underline=True; doc.add_paragraph('服务器端通过环境变量保存 Gemini API Key，前端不会直接暴露密钥。服务监听平台提供的 PORT，并支持 HTTPS。Render 免费实例长时间无访问后会休眠，首次访问可能需要等待约一分钟。')
doc.add_heading('当前限制',1)
for x in ['服务器会话和限流记录保存在进程内存中，重启或实例休眠后可能丢失。','公开访问会消耗站长配置的 Gemini API 配额，建议使用邀请制、访问密码或更严格的每日限额。','品牌入口目前是京东、淘宝和官网搜索链接，不等同于实时商品详情页。','价格、库存、尺寸和性能等外部信息若未接入实时数据源，会明确标为待核实。','当前适合演示和小规模试用，尚未包含用户登录、多用户历史记录和数据库持久化。']: doc.add_paragraph(x,style='List Bullet')
doc.add_heading('后续建议',1)
for x in ['接入实时搜索或电商数据源，补充真实产品图片、价格、库存、品牌资料和稳定详情链接。','增加自动重连、统一的错误提示和服务唤醒提示，降低 Render 免费实例休眠带来的困惑。','加入访问密码、用户级配额、每日总额度和用量监控，控制公开分享后的 API 成本。','将会话、生成页面和用户操作迁移到数据库或 Redis，支持长期保存和多用户使用。','建立跨领域评测，持续检查目标理解、约束保留、事实可靠性、计算正确性和页面适配度。']: doc.add_paragraph(x,style='List Number')
p=doc.add_paragraph('— 文档结束 —',style='Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.save(OUT); print(OUT)
